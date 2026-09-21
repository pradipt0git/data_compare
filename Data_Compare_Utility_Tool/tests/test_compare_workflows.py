"""End-to-end coverage for the Data Compare Utility workflows."""

import io
import json
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

from app import create_app
from app.compare import RED_FILL, YELLOW_FILL, detect_file_type


def workbook_bytes(sheets):
    """sheets: [(title, [[row], ...]), ...] -> in-memory .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


HEADERS = ["Account ID", "Customer Name", "Amount"]
RIGHT_HEADERS = ["AccountID", "Customer  Name", "Amount"]


class WorkflowTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="data_compare_test_")
        self.app = create_app(db_path=os.path.join(self.tmp, "compare.db"),
                              output_dir=os.path.join(self.tmp, "output"))
        self.app.config["TESTING"] = True
        self.client = self.app.test_client()

    def tearDown(self):
        self.client = None
        shutil.rmtree(self.tmp, ignore_errors=True)

    # -- helpers ---------------------------------------------------------
    def upload(self, endpoint, left, right, left_name="left.xlsx", right_name="right.xlsx"):
        return self.client.post(
            endpoint,
            data={
                "left_file": (io.BytesIO(left), left_name),
                "right_file": (io.BytesIO(right), right_name),
                "operation_id": "test-op",
            },
            content_type="multipart/form-data",
        )

    def simple_pair(self, left_rows, right_rows):
        left = workbook_bytes([("Left", [HEADERS] + left_rows)])
        right = workbook_bytes([("Right", [RIGHT_HEADERS] + right_rows)])
        return left, right

    def advanced_run(self, left, right, primary_keys):
        response = self.upload("/api/executions", left, right)
        self.assertEqual(response.status_code, 200)
        execution_id = response.get_json()["execution_id"]

        sheets = self.client.post(
            "/api/executions/%s/sheets" % execution_id,
            json={"left_sheet": "Left", "right_sheet": "Right"})
        self.assertEqual(sheets.status_code, 200)

        mapping = self.client.post(
            "/api/executions/%s/mapping" % execution_id,
            json={"primary_keys": primary_keys})
        self.assertEqual(mapping.status_code, 200)

        compare = self.client.post(
            "/api/executions/%s/compare" % execution_id,
            json={"operation_id": "test-op"})
        self.assertEqual(compare.status_code, 200, compare.get_data(as_text=True))
        return execution_id, compare.get_json()

    # -- basic workflow --------------------------------------------------
    def test_basic_selects_first_sheet_in_multi_sheet_workbooks(self):
        left = workbook_bytes([
            ("FirstLeft", [HEADERS, ["A-1", "Acme", "10"]]),
            ("SecondLeft", [HEADERS, ["Z-9", "Other", "99"]]),
        ])
        right = workbook_bytes([
            ("FirstRight", [RIGHT_HEADERS, ["A-1", "Acme", "10"]]),
            ("SecondRight", [RIGHT_HEADERS, ["Z-9", "Other", "99"]]),
        ])
        body = self.upload("/api/executions/basic", left, right).get_json()
        self.assertEqual(body["left_sheet"], "FirstLeft")
        self.assertEqual(body["right_sheet"], "FirstRight")

    def test_basic_uses_all_matched_columns_as_key_with_default_rules(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        body = self.upload("/api/executions/basic", left, right).get_json()
        self.assertEqual(body["primary_keys"], HEADERS)
        self.assertEqual(body["rules"], {
            "trim_spaces": True,
            "case_insensitive": True,
            "blank_null_equal": True,
            "normalize_numbers": True,
            "normalize_dates": True,
            "fuzzy_threshold": 0.9,
            "ignore_unmapped_columns": True,
        })

    def test_basic_completes_and_returns_compare_data(self):
        left, right = self.simple_pair(
            [["A-1", "Acme", "10"], ["A-2", "Beta", "20"]],
            [["A-1", "Acme", "10"], ["A-2", "Beta", "20"]])
        body = self.upload("/api/executions/basic", left, right).get_json()
        self.assertEqual(body["current_step"], 7)
        self.assertEqual(body["summary"]["common_count"], 2)

        results = self.client.get("/api/executions/%s/results" % body["execution_id"]).get_json()
        for key in ("common", "mismatch_first", "mismatch_second", "duplicate", "compare"):
            self.assertIn(key, results)
        self.assertEqual(len(results["compare"]), 2)
        self.assertTrue(all(row["left"] and row["right"] for row in results["compare"]))

    def test_text_files_compare_line_by_line(self):
        left = b"alpha\nbeta\ngamma\n"
        right = b"alpha\nbeta changed\ngamma\ndelta\n"
        body = self.upload("/api/executions/basic", left, right,
                           left_name="left.txt", right_name="right.txt").get_json()
        # A text file becomes a line number plus the line, keyed on the line number.
        self.assertEqual(body["left_columns"], ["Line", "Text"])
        self.assertEqual(body["primary_keys"], ["Line"])

        results = self.client.get(
            "/api/executions/%s/results" % body["execution_id"]).get_json()
        self.assertEqual(sorted(r["key_norm"] for r in results["common"]), ["1", "3"])
        changed = [r for r in results["mismatch_first"] if not r["is_missing"]]
        self.assertEqual([r["key_norm"] for r in changed], ["2"])
        self.assertEqual(changed[0]["row"]["Text"], "beta")
        only_right = [r for r in results["mismatch_second"] if r["is_missing"]]
        self.assertEqual([r["key_norm"] for r in only_right], ["4"])
        self.assertEqual(only_right[0]["row"]["Text"], "delta")

    def test_json_ignores_formatting_and_key_order(self):
        left = b'{\n  "a": 1,\n  "b": 2\n}\n'
        right = b'{"b": 2, "a": 1}'          # same data, minified, keys reordered
        body = self.upload("/api/executions/basic", left, right,
                           left_name="l.json", right_name="r.json").get_json()
        self.assertEqual(body["left_columns"], ["Line", "Text"])
        self.assertEqual(body["primary_keys"], ["Line"])

        results = self.client.get(
            "/api/executions/%s/results" % body["execution_id"]).get_json()
        self.assertEqual(len(results["common"]), 4)
        self.assertEqual(results["mismatch_first"], [])
        self.assertEqual(results["mismatch_second"], [])

    def test_json_value_change_is_reported_on_its_line(self):
        body = self.upload("/api/executions/basic", b'{"a": 1, "b": 2}', b'{"a": 1, "b": 3}',
                           left_name="l.json", right_name="r.json").get_json()
        results = self.client.get(
            "/api/executions/%s/results" % body["execution_id"]).get_json()
        changed = [r for r in results["mismatch_first"] if not r["is_missing"]]
        self.assertEqual([r["key_norm"] for r in changed], ["3"])
        self.assertEqual(changed[0]["row"]["Text"], '  "b": 2')

    def test_unparseable_json_falls_back_to_plain_lines(self):
        body = self.upload("/api/executions/basic", b"not json\nsecond line\n",
                           b"not json\nchanged line\n",
                           left_name="l.json", right_name="r.json").get_json()
        results = self.client.get(
            "/api/executions/%s/results" % body["execution_id"]).get_json()
        self.assertEqual([r["key_norm"] for r in results["common"]], ["1"])
        changed = [r for r in results["mismatch_first"] if not r["is_missing"]]
        self.assertEqual([r["key_norm"] for r in changed], ["2"])

    def test_file_types_are_detected_by_extension(self):
        self.assertEqual(detect_file_type("book.xlsx"), "excel")
        self.assertEqual(detect_file_type("data.csv"), "csv")
        self.assertEqual(detect_file_type("notes.txt"), "text")
        self.assertEqual(detect_file_type("config.json"), "json")
        with self.assertRaises(ValueError):
            detect_file_type("notes.docx")

    def test_strict_all_column_key_turns_a_change_into_missing_rows(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "11"]])
        body = self.upload("/api/executions/basic", left, right).get_json()
        summary = body["summary"]
        self.assertEqual(summary["common_count"], 0)
        self.assertEqual(summary["changed_count"], 0)
        self.assertEqual(summary["mismatch_first_count"], 1)
        self.assertEqual(summary["mismatch_second_count"], 1)

        results = self.client.get("/api/executions/%s/results" % body["execution_id"]).get_json()
        self.assertTrue(results["mismatch_first"][0]["is_missing"])
        self.assertTrue(results["mismatch_second"][0]["is_missing"])

    # -- advanced workflow -----------------------------------------------
    def test_advanced_upload_endpoint_remains_compatible(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        response = self.upload("/api/executions", left, right)
        body = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertIn("execution_id", body)
        self.assertEqual(body["left_sheets"], ["Left"])
        self.assertEqual(body["right_sheets"], ["Right"])
        self.assertEqual(body["execution"]["status"], "uploaded")

    def test_rapid_requests_receive_distinct_microsecond_ids(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        ids = set()
        for _ in range(3):
            ids.add(self.upload("/api/executions", left, right).get_json()["execution_id"])
            ids.add(self.upload("/api/executions/basic", left, right).get_json()["execution_id"])
        self.assertEqual(len(ids), 6)

    def test_advanced_compare_reports_changed_columns(self):
        left, right = self.simple_pair(
            [["A-1", "Acme", "10"], ["A-2", "Beta", "20"]],
            [["A-1", "Acme", "11"], ["A-2", "Beta", "20"]])
        execution_id, body = self.advanced_run(left, right, ["Account ID"])
        self.assertEqual(body["summary"]["changed_count"], 1)
        self.assertEqual(body["summary"]["common_count"], 1)

        results = self.client.get("/api/executions/%s/results" % execution_id).get_json()
        changed = [row for row in results["compare"] if row["mismatch_cols"]]
        self.assertEqual(changed[0]["mismatch_cols"], ["Amount"])

    def test_validation_blocks_compare_when_a_key_is_missing_on_the_right(self):
        left = workbook_bytes([("Left", [["Account ID", "Extra"], ["A-1", "x"]])])
        right = workbook_bytes([("Right", [["AccountID"], ["A-1"]])])
        response = self.upload("/api/executions", left, right)
        execution_id = response.get_json()["execution_id"]
        self.client.post("/api/executions/%s/sheets" % execution_id,
                         json={"left_sheet": "Left", "right_sheet": "Right"})
        self.client.post("/api/executions/%s/mapping" % execution_id,
                         json={"primary_keys": ["Account ID", "Extra"]})
        compare = self.client.post("/api/executions/%s/compare" % execution_id, json={})
        self.assertEqual(compare.status_code, 400)
        self.assertFalse(compare.get_json()["validation"]["ok"])

    # -- aligned compare data --------------------------------------------
    def test_duplicate_occurrences_pair_by_source_order(self):
        left, right = self.simple_pair(
            [["A-1", "First", "10"], ["A-1", "Second", "20"]],
            [["A-1", "First", "10"], ["A-1", "Second", "20"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        results = self.client.get("/api/executions/%s/results" % execution_id).get_json()
        rows = [row for row in results["compare"] if row["key_norm"] == "a-1"]
        self.assertEqual([row["occurrence"] for row in rows], [1, 2])
        self.assertEqual(rows[0]["left"]["Customer Name"], "First")
        self.assertEqual(rows[0]["right"]["Customer Name"], "First")
        self.assertEqual(rows[1]["left"]["Customer Name"], "Second")

        duplicates = results["duplicate"]
        self.assertEqual(sorted(d["side"] for d in duplicates), ["left", "right"])
        self.assertTrue(all(d["duplicate_count"] == 2 for d in duplicates))

    def test_excess_duplicate_occurrence_gets_a_null_opposite_side(self):
        left, right = self.simple_pair(
            [["A-1", "First", "10"], ["A-1", "Second", "20"]],
            [["A-1", "First", "10"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        results = self.client.get("/api/executions/%s/results" % execution_id).get_json()
        rows = [row for row in results["compare"] if row["key_norm"] == "a-1"]
        self.assertEqual(len(rows), 2)
        self.assertIsNone(rows[1]["right"])
        self.assertTrue(rows[1]["is_missing_right"])

    def test_populated_keys_sort_before_blank_keys(self):
        left, right = self.simple_pair(
            [["", "Blank one", "1"], ["A-2", "Populated", "2"], ["", "Blank two", "3"]],
            [["", "Blank one", "1"], ["A-2", "Populated", "2"], ["", "Blank two", "3"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        results = self.client.get("/api/executions/%s/results" % execution_id).get_json()
        keys = [row["key_norm"] for row in results["compare"]]
        self.assertEqual(keys[0], "a-2")
        self.assertEqual(keys[1:], ["", ""])

    # -- exports ---------------------------------------------------------
    def test_legacy_workbook_sheets_retain_fills_and_structure(self):
        left, right = self.simple_pair(
            [["A-1", "Acme", "10"], ["A-2", "Beta", "20"], ["A-3", "Only left", "30"]],
            [["A-1", "Acme", "11"], ["A-2", "Beta", "20"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        response = self.client.get("/api/executions/%s/download/excel" % execution_id)
        self.assertEqual(response.status_code, 200)

        wb = load_workbook(io.BytesIO(response.data))
        try:
            self.assertEqual(wb.sheetnames, [
                "Common data", "Data mismatch in First file",
                "Data mismatch in Second file", "Duplicate entries", "Compare"])
            first = wb["Data mismatch in First file"]
            self.assertEqual(first["A1"].value, "Key")
            changed = [r for r in first.iter_rows(min_row=2) if r[0].value == "a-1"][0]
            self.assertEqual(changed[3].fill.start_color.rgb, YELLOW_FILL.start_color.rgb)
            missing = [r for r in first.iter_rows(min_row=2) if r[0].value == "a-3"][0]
            self.assertEqual(missing[1].fill.start_color.rgb, RED_FILL.start_color.rgb)
        finally:
            wb.close()

    def test_live_compare_export_respects_columns_order_and_fills(self):
        left, right = self.simple_pair(
            [["A-1", "Acme", "10"], ["A-2", "Beta", "20"], ["A-3", "Only left", "30"]],
            [["A-1", "Acme", "11"], ["A-2", "Beta", "20"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        payload = {
            "selected_columns": ["Account ID", "Amount"],
            "row_order": [
                {"key_norm": "a-3", "occurrence": 1},
                {"key_norm": "a-1", "occurrence": 1},
            ],
        }
        response = self.client.post(
            "/api/executions/%s/download/excel/compare" % execution_id, json=payload)
        self.assertEqual(response.status_code, 200)

        wb = load_workbook(io.BytesIO(response.data))
        try:
            ws = wb["Compare"]
            header = [cell.value for cell in ws[1]]
            self.assertEqual(header, ["Account ID", "Amount", None, "AccountID", "Amount"])
            keys = [ws.cell(row=index, column=1).value for index in (2, 3)]
            self.assertEqual(keys, ["A-3", "A-1"])
            # Missing right side is red on every right cell of the first data row.
            self.assertEqual(ws.cell(row=2, column=4).fill.start_color.rgb,
                             RED_FILL.start_color.rgb)
            self.assertIsNone(ws.cell(row=2, column=3).value)
            # The changed Amount is yellow on both sides.
            self.assertEqual(ws.cell(row=3, column=2).fill.start_color.rgb,
                             YELLOW_FILL.start_color.rgb)
            self.assertEqual(ws.cell(row=3, column=5).fill.start_color.rgb,
                             YELLOW_FILL.start_color.rgb)
        finally:
            wb.close()

    def test_csv_export_and_progress_endpoints(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        execution_id, _ = self.advanced_run(left, right, ["Account ID"])
        response = self.client.get("/api/executions/%s/download/csv/common" % execution_id)
        self.assertEqual(response.status_code, 200)
        self.assertIn("Account ID", response.get_data(as_text=True))

        unknown = self.client.get("/api/progress/does-not-exist").get_json()
        self.assertEqual(unknown, {"percent": 0, "status": "Waiting to start"})

    def test_history_delete_removes_rows_and_output_folder(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        execution_id = self.upload("/api/executions/basic", left, right).get_json()["execution_id"]
        folder = os.path.join(self.app.config["OUTPUT_DIR"], execution_id)
        self.assertTrue(os.path.isdir(folder))

        self.client.delete("/api/executions/%s" % execution_id)
        self.assertFalse(os.path.exists(folder))
        listing = self.client.get("/api/executions").get_json()["executions"]
        self.assertEqual([e["execution_id"] for e in listing], [])

    def test_csv_source_files_are_supported(self):
        left = b"Account ID,Customer Name,Amount\nA-1,Acme,10\n"
        right = b"AccountID,Customer  Name,Amount\nA-1,Acme,10\n"
        response = self.client.post(
            "/api/executions/basic",
            data={"left_file": (io.BytesIO(left), "left.csv"),
                  "right_file": (io.BytesIO(right), "right.csv")},
            content_type="multipart/form-data")
        body = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(body["left_sheet"], "CSV")
        self.assertEqual(body["summary"]["common_count"], 1)

    def test_clear_db_empties_history(self):
        left, right = self.simple_pair([["A-1", "Acme", "10"]], [["A-1", "Acme", "10"]])
        self.upload("/api/executions/basic", left, right)
        self.client.post("/api/clear-db")
        listing = self.client.get("/api/executions").get_json()["executions"]
        self.assertEqual(listing, [])


if __name__ == "__main__":
    unittest.main()
