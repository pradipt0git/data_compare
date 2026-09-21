"""File reading, normalization, column matching and the comparison engine."""

import csv
import difflib
import io
import json
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

CSV_SHEET_NAME = "CSV"
TEXT_SHEET_NAME = "TEXT"
JSON_SHEET_NAME = "JSON"
# Formats with no columns of their own: read as numbered lines and compared line by line.
LINE_BASED_TYPES = ("text", "json")
# A plain text file has no headers, so it is given a line number and its line content.
# Keying on the line number is what makes the comparison run line by line.
TEXT_LINE_COLUMN = "Line"
TEXT_BODY_COLUMN = "Text"
BATCH_SIZE = 2000

DEFAULT_RULES = {
    "trim_spaces": True,
    "case_insensitive": True,
    "blank_null_equal": True,
    "normalize_numbers": True,
    "normalize_dates": True,
    "fuzzy_threshold": 0.9,
    "ignore_unmapped_columns": True,
}

DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y")

LEGACY_SHEETS = (
    ("common", "Common data"),
    ("mismatch_first", "Data mismatch in First file"),
    ("mismatch_second", "Data mismatch in Second file"),
    ("duplicate", "Duplicate entries"),
)


def default_rules():
    return dict(DEFAULT_RULES)


def normalize_header(name):
    """Headers compare with all whitespace removed and lowercased."""
    if name is None:
        return ""
    return "".join(str(name).split()).lower()


def detect_file_type(file_name):
    ext = os.path.splitext(file_name or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return "excel"
    if ext == ".csv":
        return "csv"
    if ext in (".txt", ".log"):
        return "text"
    if ext == ".json":
        return "json"
    raise ValueError(
        "Unsupported file type '%s'. Use .xlsx, .xlsm, .csv, .txt or .json."
        % (ext or "unknown"))


def _format_number(value):
    """Render numbers without insignificant trailing zeroes."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    try:
        dec = Decimal(str(value)).normalize()
    except (InvalidOperation, ValueError):
        return str(value)
    text = format(dec, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def cell_to_text(value):
    """Convert a raw cell value into the canonical string kept in SQLite."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _format_number(value)
    if isinstance(value, Decimal):
        return _format_number(value)
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value)


def list_sheets(file_path, file_type):
    if file_type == "csv":
        return [CSV_SHEET_NAME]
    if file_type == "text":
        return [TEXT_SHEET_NAME]
    if file_type == "json":
        return [JSON_SHEET_NAME]
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def iter_rows(file_path, file_type, sheet_name):
    """Yield lists of text cells, header row first."""
    if file_type in LINE_BASED_TYPES:
        # Synthesised header, then one row per line: the line number and the line itself.
        yield [TEXT_LINE_COLUMN, TEXT_BODY_COLUMN]
        for number, line in enumerate(_read_lines(file_path, file_type), start=1):
            yield [str(number), line]
        return

    if file_type == "csv":
        with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                yield [cell_to_text(cell) for cell in row]
        return

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            yield [cell_to_text(cell) for cell in row]
    finally:
        wb.close()


def _read_lines(file_path, file_type):
    """The lines of a line-based file.

    JSON is re-serialised with sorted keys and a fixed indent before it is split, so that
    two files holding the same data do not read as different merely because one is
    minified, indented differently, or lists its keys in another order. A file that does
    not parse is compared as plain text instead.
    """
    with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
        raw = handle.read()
    if file_type == "json":
        try:
            parsed = json.loads(raw)
        except ValueError:
            pass
        else:
            return json.dumps(parsed, indent=2, sort_keys=True,
                              ensure_ascii=False).split("\n")
    return raw.splitlines()


def read_headers(file_path, file_type, sheet_name):
    for row in iter_rows(file_path, file_type, sheet_name):
        headers = []
        seen = {}
        for index, cell in enumerate(row):
            name = (cell or "").strip() or "Column%d" % (index + 1)
            if name in seen:
                seen[name] += 1
                name = "%s_%d" % (name, seen[name])
            else:
                seen[name] = 0
            headers.append(name)
        return headers
    return []


def normalize_value(value, rules=None):
    """Apply the configured rules to produce a comparable string."""
    rules = rules or DEFAULT_RULES
    if value is None:
        return ""
    text = value if isinstance(value, str) else cell_to_text(value)

    if rules.get("trim_spaces", True):
        text = text.strip()

    if rules.get("blank_null_equal", True) and text.strip().lower() in ("", "null", "none", "nan"):
        return ""

    if rules.get("normalize_numbers", True):
        candidate = text.replace(",", "").strip()
        try:
            dec = Decimal(candidate)
        except (InvalidOperation, ValueError):
            pass
        else:
            if candidate not in ("", "-", "+"):
                text = _format_number(dec)
                return text.lower() if rules.get("case_insensitive", True) else text

    if rules.get("normalize_dates", True):
        stamp = text.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(stamp, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue

    if rules.get("case_insensitive", True):
        text = text.lower()
    return text


def auto_match_columns(left_columns, right_columns, rules=None):
    """Exact normalized matches first, then difflib fuzzy matching."""
    rules = rules or DEFAULT_RULES
    threshold = float(rules.get("fuzzy_threshold", 0.9))

    right_norm = [normalize_header(name) for name in right_columns]
    used_right = set()
    matches = [None] * len(left_columns)

    for li, left_name in enumerate(left_columns):
        target = normalize_header(left_name)
        for ri, candidate in enumerate(right_norm):
            if ri in used_right or candidate != target:
                continue
            used_right.add(ri)
            matches[li] = {
                "left_index": li,
                "left_name": left_name,
                "right_index": ri,
                "right_name": right_columns[ri],
                "status": "matched",
                "score": 1.0,
            }
            break

    for li, left_name in enumerate(left_columns):
        if matches[li] is not None:
            continue
        target = normalize_header(left_name)
        best_index, best_score = None, 0.0
        for ri, candidate in enumerate(right_norm):
            if ri in used_right:
                continue
            score = difflib.SequenceMatcher(None, target, candidate).ratio()
            if score > best_score:
                best_index, best_score = ri, score
        if best_index is not None and best_score >= threshold:
            used_right.add(best_index)
            matches[li] = {
                "left_index": li,
                "left_name": left_name,
                "right_index": best_index,
                "right_name": right_columns[best_index],
                "status": "partial",
                "score": round(best_score, 4),
            }
        else:
            matches[li] = {
                "left_index": li,
                "left_name": left_name,
                "right_index": None,
                "right_name": None,
                "status": "missing",
                "score": round(best_score, 4) if best_index is not None else 0.0,
            }

    return matches


def matches_from_pairs(pairs):
    """Build matches from the rows as arranged on the mapping screen.

    Each pair is [left_name, right_name]; either side may be blank, which is how the
    screen represents a column with no counterpart. Position is the whole truth here,
    so a name the user dragged opposite another one is paired with it even when the
    two names look nothing alike.
    """
    matches = []
    for pair in pairs or []:
        left_name = (pair or [None, None])[0]
        right_name = (pair or [None, None])[1] if len(pair or []) > 1 else None
        if not left_name:
            continue
        index = len(matches)
        matches.append({
            "left_index": index,
            "left_name": left_name,
            "right_index": index if right_name else None,
            "right_name": right_name or None,
            "status": "matched" if right_name else "missing",
            "score": 1.0 if right_name else 0.0,
        })
    return matches


def resolve_column_matches(matches):
    """Map names in both directions and list the comparable left columns."""
    right_name_by_left = {}
    left_name_by_right = {}
    comparable_left_columns = []
    for match in matches or []:
        if match.get("status") not in ("matched", "partial"):
            continue
        left_name = match.get("left_name")
        right_name = match.get("right_name")
        if not left_name or not right_name:
            continue
        right_name_by_left[left_name] = right_name
        left_name_by_right[right_name] = left_name
        comparable_left_columns.append(left_name)
    return {
        "right_name_by_left": right_name_by_left,
        "left_name_by_right": left_name_by_right,
        "comparable_left_columns": comparable_left_columns,
    }


# ---------------------------------------------------------------------------
# Loading rows into SQLite
# ---------------------------------------------------------------------------

def _source_index_map(headers, columns, matches, side):
    """Map canonical left column names to a source column index for `side`."""
    header_norm = [normalize_header(name) for name in headers]
    index_by_norm = {}
    for index, norm in enumerate(header_norm):
        index_by_norm.setdefault(norm, index)

    resolved = resolve_column_matches(matches)
    right_name_by_left = resolved["right_name_by_left"]

    source_index = {}
    for left_name in columns:
        wanted = left_name if side == "left" else right_name_by_left.get(left_name)
        if wanted is None:
            continue
        index = index_by_norm.get(normalize_header(wanted))
        if index is not None:
            source_index[left_name] = index
    return source_index


def build_key_norm(payload, primary_keys, rules):
    return "|".join(normalize_value(payload.get(key, ""), rules) for key in primary_keys)


def stream_to_sqlite(conn, execution_id, side, file_path, file_type, sheet_name,
                     columns, matches, primary_keys, rules, progress=None):
    """Stream one side of the comparison into data_rows under canonical keys."""
    rules = rules or DEFAULT_RULES
    conn.execute(
        "DELETE FROM data_rows WHERE execution_id = ? AND side = ?",
        (execution_id, side),
    )

    stream = iter_rows(file_path, file_type, sheet_name)
    try:
        headers = next(stream)
    except StopIteration:
        return 0

    header_names = read_headers(file_path, file_type, sheet_name)
    source_index = _source_index_map(header_names or headers, columns, matches, side)

    batch = []
    row_index = 0
    for raw in stream:
        if not any((cell or "").strip() for cell in raw):
            continue
        payload = {}
        for column in columns:
            index = source_index.get(column)
            payload[column] = raw[index] if index is not None and index < len(raw) else ""
        key_norm = build_key_norm(payload, primary_keys, rules)
        batch.append((execution_id, side, row_index, key_norm, json.dumps(payload)))
        row_index += 1
        if len(batch) >= BATCH_SIZE:
            conn.executemany(
                "INSERT INTO data_rows (execution_id, side, row_index, key_norm, row_json)"
                " VALUES (?, ?, ?, ?, ?)",
                batch,
            )
            batch = []
            if progress:
                progress(row_index)

    if batch:
        conn.executemany(
            "INSERT INTO data_rows (execution_id, side, row_index, key_norm, row_json)"
            " VALUES (?, ?, ?, ?, ?)",
            batch,
        )
    return row_index


def _load_side(conn, execution_id, side):
    cursor = conn.execute(
        "SELECT row_index, key_norm, row_json FROM data_rows"
        " WHERE execution_id = ? AND side = ? ORDER BY row_index",
        (execution_id, side),
    )
    grouped = {}
    for row in cursor.fetchall():
        grouped.setdefault(row["key_norm"], []).append(
            {"row_index": row["row_index"], "row": json.loads(row["row_json"])}
        )
    return grouped


# ---------------------------------------------------------------------------
# Comparison
# ---------------------------------------------------------------------------

def run_comparison(conn, execution_id, columns, matches, rules, output_dir,
                   progress=None):
    """Compare both persisted sides and store result_rows plus a legacy workbook."""
    rules = rules or DEFAULT_RULES
    resolved = resolve_column_matches(matches)
    comparable = resolved["comparable_left_columns"]

    conn.execute("DELETE FROM result_rows WHERE execution_id = ?", (execution_id,))

    left = _load_side(conn, execution_id, "left")
    right = _load_side(conn, execution_id, "right")

    duplicate_keys = set()
    duplicate_rows = []
    for side, grouped in (("left", left), ("right", right)):
        for key, occurrences in grouped.items():
            if key and len(occurrences) > 1:
                duplicate_keys.add(key)
                duplicate_rows.append((
                    execution_id, "duplicate", side, key,
                    json.dumps(occurrences[0]["row"]), json.dumps([]),
                    len(occurrences), 0,
                ))

    inserts = list(duplicate_rows)
    common_count = 0
    changed_count = 0
    mismatch_first_count = 0
    mismatch_second_count = 0

    all_keys = set(left) | set(right)
    for key in all_keys:
        if key in duplicate_keys:
            continue
        left_rows = left.get(key, [])
        right_rows = right.get(key, [])
        pairs = max(len(left_rows), len(right_rows))
        for position in range(pairs):
            left_entry = left_rows[position] if position < len(left_rows) else None
            right_entry = right_rows[position] if position < len(right_rows) else None

            if left_entry and not right_entry:
                inserts.append((execution_id, "mismatch_first", "left", key,
                                json.dumps(left_entry["row"]), json.dumps([]), None, 1))
                mismatch_first_count += 1
                continue
            if right_entry and not left_entry:
                inserts.append((execution_id, "mismatch_second", "right", key,
                                json.dumps(right_entry["row"]), json.dumps([]), None, 1))
                mismatch_second_count += 1
                continue

            mismatch_cols = [
                column for column in comparable
                if normalize_value(left_entry["row"].get(column, ""), rules)
                != normalize_value(right_entry["row"].get(column, ""), rules)
            ]
            if mismatch_cols:
                payload = json.dumps(mismatch_cols)
                inserts.append((execution_id, "mismatch_first", "left", key,
                                json.dumps(left_entry["row"]), payload, None, 0))
                inserts.append((execution_id, "mismatch_second", "right", key,
                                json.dumps(right_entry["row"]), payload, None, 0))
                mismatch_first_count += 1
                mismatch_second_count += 1
                changed_count += 1
            else:
                inserts.append((execution_id, "common", "left", key,
                                json.dumps(left_entry["row"]), json.dumps([]), None, 0))
                common_count += 1

    for start in range(0, len(inserts), BATCH_SIZE):
        conn.executemany(
            "INSERT INTO result_rows (execution_id, result_type, side, key_norm, row_json,"
            " mismatch_cols_json, duplicate_count, is_missing) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            inserts[start:start + BATCH_SIZE],
        )

    if progress:
        progress("workbook")

    # The heavy side-by-side Compare sheet is deferred to the download endpoint.
    workbook_bytes = export_workbook(conn, execution_id, columns, matches, rules,
                                     include_compare=False)
    os.makedirs(output_dir, exist_ok=True)
    output_excel = os.path.join(output_dir, "comparison_%s.xlsx" % execution_id)
    with open(output_excel, "wb") as handle:
        handle.write(workbook_bytes)

    return {
        "common_count": common_count,
        "mismatch_first_count": mismatch_first_count,
        "mismatch_second_count": mismatch_second_count,
        "duplicate_count": len(duplicate_rows),
        "output_excel": output_excel,
        "changed_count": changed_count,
        "duplicate_key_count": len(duplicate_keys),
    }


# ---------------------------------------------------------------------------
# Aligned compare data
# ---------------------------------------------------------------------------

def _project(row, columns):
    return {column: row.get(column, "") for column in columns}


def build_aligned_rows(conn, execution_id, columns, rules):
    """Pair both persisted sides by key and occurrence without re-reading files."""
    rules = rules or DEFAULT_RULES
    left = _load_side(conn, execution_id, "left")
    right = _load_side(conn, execution_id, "right")

    aligned = []
    for key in set(left) | set(right):
        left_rows = left.get(key, [])
        right_rows = right.get(key, [])
        for position in range(max(len(left_rows), len(right_rows))):
            left_entry = left_rows[position] if position < len(left_rows) else None
            right_entry = right_rows[position] if position < len(right_rows) else None
            left_row = _project(left_entry["row"], columns) if left_entry else None
            right_row = _project(right_entry["row"], columns) if right_entry else None

            mismatch_cols = []
            if left_row is not None and right_row is not None:
                mismatch_cols = [
                    column for column in columns
                    if normalize_value(left_row.get(column, ""), rules)
                    != normalize_value(right_row.get(column, ""), rules)
                ]

            aligned.append({
                "key_norm": key,
                "occurrence": position + 1,
                "left_row_index": left_entry["row_index"] if left_entry else None,
                "right_row_index": right_entry["row_index"] if right_entry else None,
                "left": left_row,
                "right": right_row,
                "mismatch_cols": mismatch_cols,
                "is_missing_left": left_row is None,
                "is_missing_right": right_row is None,
            })

    # Blank keys sort after every populated key; occurrence keeps order stable.
    aligned.sort(key=lambda item: (item["key_norm"] == "", item["key_norm"], item["occurrence"]))
    return aligned


def row_identity(row):
    return "%s||%s" % (row.get("key_norm", ""), row.get("occurrence", 1))


def order_aligned_rows(aligned, row_order):
    """Reorder aligned rows to mirror the client's filtered/sorted Compare view."""
    if not row_order:
        return aligned
    wanted = []
    for item in row_order:
        if isinstance(item, dict):
            wanted.append("%s||%s" % (item.get("key_norm", ""), item.get("occurrence", 1)))
        else:
            wanted.append(str(item))
    by_identity = {row_identity(row): row for row in aligned}
    ordered = [by_identity[identity] for identity in wanted if identity in by_identity]
    return ordered


# ---------------------------------------------------------------------------
# Workbook export
# ---------------------------------------------------------------------------

def _write_legacy_sheet(ws, columns, rows):
    ws.append(["Key"] + list(columns) + ["Info"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for record in rows:
        payload = record["row"]
        mismatch_cols = record.get("mismatch_cols") or []
        info = ""
        if record.get("is_missing"):
            info = "Missing on the other side"
        elif record.get("duplicate_count"):
            info = "Duplicate key x%d" % record["duplicate_count"]
        elif mismatch_cols:
            info = "Differs: %s" % ", ".join(mismatch_cols)
        ws.append([record.get("key_norm", "")] + [payload.get(c, "") for c in columns] + [info])
        written = ws[ws.max_row]
        for offset, column in enumerate(columns, start=1):
            if record.get("is_missing"):
                written[offset].fill = RED_FILL
            elif column in mismatch_cols:
                written[offset].fill = YELLOW_FILL


def _fetch_results(conn, execution_id, result_type):
    cursor = conn.execute(
        "SELECT side, key_norm, row_json, mismatch_cols_json, duplicate_count, is_missing"
        " FROM result_rows WHERE execution_id = ? AND result_type = ? ORDER BY id",
        (execution_id, result_type),
    )
    records = []
    for row in cursor.fetchall():
        records.append({
            "side": row["side"],
            "key_norm": row["key_norm"],
            "row": json.loads(row["row_json"]),
            "mismatch_cols": json.loads(row["mismatch_cols_json"] or "[]"),
            "duplicate_count": row["duplicate_count"],
            "is_missing": bool(row["is_missing"]),
        })
    return records


def _write_compare_sheet(ws, left_columns, right_labels, aligned):
    header = list(left_columns) + [""] + list(right_labels)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
        if cell.value:
            cell.fill = HEADER_FILL

    separator = len(left_columns) + 1
    for record in aligned:
        left_row = record.get("left")
        right_row = record.get("right")
        mismatch_cols = set(record.get("mismatch_cols") or [])
        values = [("" if left_row is None else left_row.get(c, "")) for c in left_columns]
        values.append("")
        values.extend(("" if right_row is None else right_row.get(c, "")) for c in left_columns)
        ws.append(values)
        written = ws[ws.max_row]

        for offset, column in enumerate(left_columns):
            left_cell = written[offset]
            right_cell = written[separator + offset]
            if left_row is None:
                left_cell.fill = RED_FILL
            if right_row is None:
                right_cell.fill = RED_FILL
            if column in mismatch_cols:
                left_cell.fill = YELLOW_FILL
                right_cell.fill = YELLOW_FILL


def export_workbook(conn, execution_id, columns, matches, rules, include_compare=True,
                    selected_columns=None, row_order=None):
    """Build the legacy workbook in memory, optionally with the Compare sheet."""
    rules = rules or DEFAULT_RULES
    resolved = resolve_column_matches(matches)
    comparable = resolved["comparable_left_columns"]

    wb = Workbook()
    wb.remove(wb.active)
    for result_type, title in LEGACY_SHEETS:
        ws = wb.create_sheet(title=title)
        _write_legacy_sheet(ws, columns, _fetch_results(conn, execution_id, result_type))

    if include_compare:
        chosen = [c for c in (selected_columns or comparable) if c in comparable] or comparable
        aligned = build_aligned_rows(conn, execution_id, chosen, rules)
        aligned = order_aligned_rows(aligned, row_order)
        right_labels = [resolved["right_name_by_left"].get(c, c) for c in chosen]
        _write_compare_sheet(wb.create_sheet(title="Compare"), chosen, right_labels, aligned)

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def results_to_csv(records, columns):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Key"] + list(columns) + ["Info"])
    for record in records:
        payload = record["row"]
        info = ""
        if record.get("is_missing"):
            info = "Missing on the other side"
        elif record.get("duplicate_count"):
            info = "Duplicate key x%d" % record["duplicate_count"]
        elif record.get("mismatch_cols"):
            info = "Differs: %s" % ", ".join(record["mismatch_cols"])
        writer.writerow([record.get("key_norm", "")] + [payload.get(c, "") for c in columns] + [info])
    return buffer.getvalue()
